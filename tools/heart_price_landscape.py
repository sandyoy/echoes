# -*- coding: utf-8 -*-
"""生成 A4 横向 心脏康复套餐价格表"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = "/home/ubuntu/echoes/docs/圆爱康/03_产出方案/心脏康复_套餐价格表_v4.1_横版_20260923.docx"

doc = Document()

# ============ 页面 = A4 横向 ============
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = sec.bottom_margin = Cm(1.5)

# 中文字体
style = doc.styles["Normal"]
style.font.name = "微软雅黑"
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
style.font.size = Pt(9)

def set_cell(cell, text, bold=False, size=9, align="left", color=None, bg=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            p = cell.add_paragraph()
            p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
        run = p.add_run(line)
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(size)
        run.bold = bold
        if color:
            run.font.color.rgb = color
    if bg:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), bg)
        tcPr.append(shd)

# ============ 标题 ============
h = doc.add_paragraph()
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h.add_run("圆爱康 · 心脏康复 套餐价格表")
r.font.name = "微软雅黑"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
r.font.size = Pt(16)
r.bold = True

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("v4.1 · 2026-09-23 · 内部资料　｜　用途：销售随身带，见医院主任时用，主任在院内项目后打勾，当场形成本院价格")
r.font.name = "微软雅黑"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

# ============ 结构说明 ============
p = doc.add_paragraph()
r = p.add_run("结构说明：居家部分拆成两条独立线，可自由组合。")
r.bold = True
r.font.size = Pt(10)
r.font.name = "微软雅黑"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

p = doc.add_paragraph()
r = p.add_run("线一 超声治疗 —— 选次数：20 / 30 / 40 / 60 / 80 / 100 次　｜　"
              "线二 运动康复 —— 选计划：静养 / 起步 / 标准 / 强化 / 长期\n"
              "价格 = 超声治疗价格（按次数） ＋ 医院勾选的院内项目加价　｜　运动康复随选，不加价")
r.font.size = Pt(9)
r.font.name = "微软雅黑"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

# ============ 主表 ============
headers = ["套餐名称", "超声治疗\n（次数）", "运动康复\n（勾选一个计划）", "院内可选项目（医院勾选·勾了加价）", "价格"]

rows = [
    ["冠心病\n套餐一", "20 次",
     "☐ 静养计划\n☐ 起步计划\n☐ 标准计划\n☐ 强化计划\n☐ 长期计划",
     "☐ 六分钟步行试验　☐ 心脏彩超／射血分数　☐ 动态心电图 Holter\n☐ 血脂血糖等化验　☐ 营养体重管理处方　☐ 运动处方",
     "¥ 待定"],
    ["冠心病\n套餐二", "30 次",
     "☐ 静养计划\n☐ 起步计划\n☐ 标准计划\n☐ 强化计划\n☐ 长期计划",
     "☐ 六分钟步行试验　☐ 心脏彩超／射血分数　☐ 动态心电图 Holter\n☐ 血脂血糖等化验　☐ 营养体重管理处方　☐ 运动处方",
     "¥ 待定"],
    ["冠心病\n套餐三", "40 次",
     "☐ 静养计划\n☐ 起步计划\n☐ 标准计划\n☐ 强化计划\n☐ 长期计划",
     "☐ 心肺运动试验 CPET　☐ 六分钟步行试验　☐ 心脏彩超／射血分数\n☐ 动态心电图 Holter　☐ 血脂血糖等化验　☐ 营养体重管理处方　☐ 运动处方",
     "¥ 待定"],
    ["冠心病\n套餐四", "60 次",
     "☐ 静养计划\n☐ 起步计划\n☐ 标准计划\n☐ 强化计划\n☐ 长期计划",
     "☐ 体外反搏 EECP　☐ 心肺运动试验 CPET　☐ 六分钟步行试验\n☐ 心脏彩超／射血分数　☐ 动态心电图 Holter　☐ 运动负荷心电图\n☐ 血脂血糖等化验　☐ 营养体重管理处方　☐ 运动处方",
     "¥ 待定"],
    ["冠心病\n套餐五", "80 次",
     "☐ 静养计划\n☐ 起步计划\n☐ 标准计划\n☐ 强化计划\n☐ 长期计划",
     "☐ 体外反搏 EECP　☐ 心肺运动试验 CPET　☐ 六分钟步行试验\n☐ 心脏彩超／射血分数　☐ 动态心电图 Holter　☐ 运动负荷心电图\n☐ 血脂血糖等化验　☐ 心理状态量表　☐ 睡眠呼吸评估\n☐ 营养体重管理处方　☐ 运动处方",
     "¥ 待定"],
    ["冠心病\n套餐六", "100 次",
     "☐ 静养计划\n☐ 起步计划\n☐ 标准计划\n☐ 强化计划\n☐ 长期计划",
     "☐ 体外反搏 EECP　☐ 心肺运动试验 CPET　☐ 六分钟步行试验\n☐ 心脏彩超／射血分数　☐ 动态心电图 Holter　☐ 运动负荷心电图\n☐ 血脂血糖等化验　☐ 心理状态量表　☐ 睡眠呼吸评估\n☐ ADL 日常生活能力评估　☐ 躯体化症状评估\n☐ 营养体重管理处方　☐ 运动处方",
     "¥ 待定"],
]

t = doc.add_table(rows=1 + len(rows), cols=5)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER

# 表头
for j, htext in enumerate(headers):
    set_cell(t.rows[0].cells[j], htext, bold=True, size=9, align="center", bg="D9E2F3")

# 数据行
for i, row in enumerate(rows, start=1):
    for j, val in enumerate(row):
        set_cell(t.rows[i].cells[j], val, size=8.5,
                 align="center" if j in (0, 1, 4) else "left")

# 列宽
widths = [Cm(2.4), Cm(1.8), Cm(4.0), Cm(13.5), Cm(2.3)]
for row in t.rows:
    for idx, w in enumerate(widths):
        row.cells[idx].width = w

# 总价行
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("总价格：______________ 元　（超声治疗价格 ＋ 医院勾选院内项目加价）")
r.bold = True
r.font.size = Pt(11)
r.font.name = "微软雅黑"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

# ============ 第二页：五个计划参考 ============
doc.add_page_break()

h2 = doc.add_paragraph()
r = h2.add_run("附：五个运动康复计划（选择参考）")
r.bold = True
r.font.size = Pt(12)
r.font.name = "微软雅黑"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

plan_h = ["计划", "患者状态", "动作数", "默认每天时长", "包含动作"]
plan_rows = [
    ["静养计划", "术后当日／活动受限", "4", "5 分钟", "腹式呼吸、缩唇呼吸、坐位踝泵、运动后自测"],
    ["起步计划", "术后1-2周／日常活动有气短", "5", "10 分钟", "腹式呼吸、坐位肩颈放松、原地踏步、站立提踵、运动后自测"],
    ["标准计划", "术后2-8周／稳定冠心病", "5", "15 分钟", "腹式呼吸、坐位肩颈放松、室内慢走、坐位抬腿、运动后自测"],
    ["强化计划", "术后2-3月起／希望恢复运动能力", "9", "25 分钟", "腹式呼吸、肩颈放松、室内慢走、快走、坐位抬腿、靠墙静蹲、八段锦、太极拳、运动后自测"],
    ["长期计划", "长期维持／高血压糖尿病合并风险", "14", "30 分钟", "全套 14 个动作自由组合"],
]
t2 = doc.add_table(rows=1 + len(plan_rows), cols=5)
t2.style = "Table Grid"
for j, htext in enumerate(plan_h):
    set_cell(t2.rows[0].cells[j], htext, bold=True, size=9, align="center", bg="D9E2F3")
for i, row in enumerate(plan_rows, start=1):
    for j, val in enumerate(row):
        set_cell(t2.rows[i].cells[j], val, size=8.5, align="center" if j in (0,2,3) else "left")
w2 = [Cm(2.4), Cm(5.5), Cm(1.6), Cm(2.4), Cm(12.1)]
for row in t2.rows:
    for idx, w in enumerate(w2):
        row.cells[idx].width = w

# ============ 第三页：院内可选项目清单 ============
doc.add_page_break()
h3 = doc.add_paragraph()
r = h3.add_run("附：院内可选项目清单（医院勾选用）")
r.bold = True
r.font.size = Pt(12)
r.font.name = "微软雅黑"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

items = [
    ("评估检查类", "心肺运动试验 CPET（院内运动心肺功能评估，确定安全运动阈值）；六分钟步行试验（院内活动耐力评估）；运动负荷心电图（心肌缺血筛查）；心脏彩超／射血分数（结构功能评估）；动态心电图 Holter（心律评估）；血脂血糖等化验（危险因素评估）；心理状态量表（焦虑抑郁筛查）；睡眠呼吸评估（睡眠呼吸暂停排查）；ADL 日常生活能力评估（自理能力评估）；躯体化症状评估（躯体化量表筛查）"),
    ("院内治疗类", "体外反搏 EECP（院内体外反搏治疗，每次约半小时）；其他院内现有治疗（由医院填写）"),
    ("处方类", "运动处方（运动方案制定，与动作套组配套）；营养体重管理处方（营养与体重管理方案）"),
]
t3 = doc.add_table(rows=1 + len(items), cols=2)
t3.style = "Table Grid"
set_cell(t3.rows[0].cells[0], "类别", bold=True, size=9, align="center", bg="D9E2F3")
set_cell(t3.rows[0].cells[1], "项目内容（医院认为可一并开展的，在价格表中打勾）", bold=True, size=9, align="center", bg="D9E2F3")
for i, (cat, content) in enumerate(items, start=1):
    set_cell(t3.rows[i].cells[0], cat, bold=True, size=9, align="center")
    set_cell(t3.rows[i].cells[1], content, size=8.5, align="left")
for row in t3.rows:
    row.cells[0].width = Cm(3.0)
    row.cells[1].width = Cm(21.0)

p = doc.add_paragraph()
r = p.add_run("说明：院内项目由医院自行定价与执行，圆爱康不参与该部分定价。医院有什么勾什么，没有不勾，不强制。")
r.font.size = Pt(9)
r.font.name = "微软雅黑"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

doc.save(OUT)
print("saved:", OUT)
